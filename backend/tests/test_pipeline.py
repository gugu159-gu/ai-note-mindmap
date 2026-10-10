"""管线测试：兜底提取、Markdown 转换、鉴权、上传/复习接口。

- 全程无需 API Key（未配置时自动走本地规则兜底提取）
- 使用**独立的测试数据库**（临时目录），不会污染 backend/ai_notes.db 里的真实数据
"""
import io
import os
import sys
import tempfile

# 必须在导入 app 之前设置，决定 SQLite 文件位置
TEST_DB = os.path.join(tempfile.gettempdir(), "ai_notes_test.db")
if os.path.exists(TEST_DB):
    os.remove(TEST_DB)
os.environ["AI_NOTES_DB"] = TEST_DB

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402
from app.services.fallback import fallback_extract  # noqa: E402
from app.services.knowledge import knowledge_to_markdown, markdown_to_knowledge  # noqa: E402
from app.services.parser import parse_pptx  # noqa: E402

DEFAULT_USER = "tester"
DEFAULT_PASSWORD = "test123456"


def make_client(username: str = DEFAULT_USER, password: str = DEFAULT_PASSWORD) -> TestClient:
    """建立一个已登录的客户端（每个客户端有独立 Cookie 会话）。"""
    c = TestClient(app)
    resp = c.post("/api/auth/register", json={"username": username, "password": password})
    if resp.status_code == 400:  # 已注册过 → 改为登录
        resp = c.post("/api/auth/login", json={"username": username, "password": password})
    assert resp.status_code == 200, resp.text
    return c


client = make_client()

SAMPLE_TEXT = """
# 第一章 软件过程
软件过程是软件生命周期中的一系列活动
- 需求分析
- 设计
- 编码与测试
1.1 过程模型
瀑布模型：阶段化、文档驱动
增量模型：分批交付
# 第二章 需求工程
需求获取方法
用户访谈
原型法
"""


def test_fallback_pptx_style():
    tree = fallback_extract(SAMPLE_TEXT, "测试课件")
    assert tree["title"] == "测试课件"
    titles = [c["title"] for c in tree["children"]]
    assert "第一章 软件过程" in titles
    assert "第二章 需求工程" in titles
    # 一级标题下的编号标题成为二级节点
    ch1 = next(c for c in tree["children"] if c["title"] == "第一章 软件过程")
    assert ch1["children"] and ch1["children"][0]["title"] == "过程模型"
    assert len(ch1["key_points"]) >= 1


def test_fallback_flat_text():
    tree = fallback_extract("苹果 香蕉 橙子\n都是水果", "水果")
    assert tree["title"] == "水果"
    assert len(tree["key_points"]) >= 2


def test_fallback_empty():
    tree = fallback_extract("   \n  ", "空")
    assert tree["title"] == "空"


def test_knowledge_markdown_roundtrip():
    tree = {
        "title": "根",
        "key_points": ["要点A"],
        "children": [{"title": "子", "key_points": ["要点B"], "children": []}],
    }
    md = knowledge_to_markdown(tree)
    assert "# 根" in md
    assert "## 子" in md
    assert "- 要点A" in md
    tree2 = markdown_to_knowledge(md)
    # markdown 往返：根节点固定为「知识点」，原根成为其一级子节点
    assert tree2["children"][0]["title"] == "根"
    assert tree2["children"][0]["children"][0]["title"] == "子"


def test_pptx_parser():
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # 标题+内容
    slide.shapes.title.text = "第三章 测试"
    body = slide.placeholders[1].text_frame
    body.text = "要点一"
    body.add_paragraph().text = "要点二"
    buf = io.BytesIO()
    prs.save(buf)
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_tmp.pptx")
    with open(path, "wb") as f:
        f.write(buf.getvalue())
    text = parse_pptx(path)
    os.remove(path)
    assert "# 第三章 测试" in text
    tree = fallback_extract(text, "PPT")
    assert tree["children"][0]["title"] == "第三章 测试"


def test_upload_txt_endpoint():
    resp = client.post(
        "/api/upload",
        files={"file": ("demo.txt", SAMPLE_TEXT.encode("utf-8"), "text/plain")},
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["status"] == "done"
    assert data["outline"]
    assert data["knowledge"]["title"]
    # 任务历史可查
    tasks = client.get("/api/tasks").json()
    assert any(t["id"] == data["id"] for t in tasks)


def test_upload_bad_ext():
    resp = client.post("/api/upload", files={"file": ("x.exe", b"abc", "application/octet-stream")})
    assert resp.status_code == 400


def test_upload_empty_file():
    """空文件应被拒绝，且不产生任务记录。"""
    before = len(client.get("/api/tasks").json())
    resp = client.post("/api/upload", files={"file": ("empty.txt", b"", "text/plain")})
    assert resp.status_code == 400, resp.text
    assert len(client.get("/api/tasks").json()) == before


def test_upload_too_large():
    """超过 50MB 上限应返回 413，并清理已落盘的临时文件。"""
    from app.routers.upload import MAX_UPLOAD_BYTES, UPLOAD_DIR

    before = len(client.get("/api/tasks").json())
    before_files = len(os.listdir(UPLOAD_DIR))
    big = b"a" * (MAX_UPLOAD_BYTES + 1024)
    resp = client.post("/api/upload", files={"file": ("big.txt", big, "text/plain")})
    assert resp.status_code == 413, resp.text[:200]
    assert len(client.get("/api/tasks").json()) == before
    assert len(os.listdir(UPLOAD_DIR)) == before_files


def test_text_quality_filter():
    """乱码文本（如字形无映射的 PDF）应被判为不可用。"""
    from app.services.parser import _meaningful_ratio, _text_is_usable

    assert _text_is_usable("这是正常的课件文本内容，包含足够数量的中文字符，可以用于提取知识点结构。")
    assert not _text_is_usable("··· ·····\n1.1 OSI ····")
    assert _meaningful_ratio("··· ·····") < 0.5
    assert not _text_is_usable("太短")


def test_normalize_depth_cap():
    """LLM 返回超过 3 层时应被截断到 3 层。"""
    from app.services.extractor import normalize_knowledge

    deep = {"title": "L1", "children": [{"title": "L2", "children": [{"title": "L3", "children": [{"title": "L4"}]}]}]}
    tree = normalize_knowledge(deep)
    assert tree["title"] == "L1"
    assert tree["children"][0]["title"] == "L2"
    assert tree["children"][0]["children"][0]["title"] == "L3"
    # 第 4 层被截断
    assert tree["children"][0]["children"][0]["children"] == []


def test_normalize_points_limit():
    """要点去重、限长、每节点最多 5 条。"""
    from app.services.extractor import normalize_knowledge

    tree = normalize_knowledge(
        {"title": "T", "key_points": ["a", "a", "b", "c", "d", "e", "f", "g"]}
    )
    assert tree["key_points"] == ["a", "b", "c", "d", "e"]


def test_review_flatten_knowledge():
    """知识点树摊平：排除无要点根节点，保留带要点根节点，路径拼接正确。"""
    from app.services.review import flatten_knowledge

    tree = {
        "title": "计算机网络",
        "children": [
            {"title": "2.1 OSI", "key_points": ["物理层"], "children": [{"title": "物理层", "key_points": ["比特流"]}]}
        ],
    }
    items = flatten_knowledge(tree)
    paths = [i["path"] for i in items]
    assert paths == ["计算机网络 / 2.1 OSI", "计算机网络 / 2.1 OSI / 物理层"]

    # 根节点自身带要点时纳入复习
    tree2 = {"title": "根", "key_points": ["总述"], "children": [{"title": "子", "key_points": ["x"]}]}
    items2 = flatten_knowledge(tree2)
    assert items2[0]["title"] == "根" and items2[0]["key_points"] == ["总述"]

    # 无要点且无子节点的根 → 不产生复习项
    assert flatten_knowledge({"title": "空"}) == []


def test_review_schedule_progression():
    """艾宾浩斯间隔推进：1→2→4→7→15→30 天，走完标记已掌握。"""
    from datetime import datetime, timedelta

    from app.models import ReviewItem
    from app.services.review import EBBINGHAUS_INTERVALS, MAX_STAGE, advance, next_interval_days

    assert EBBINGHAUS_INTERVALS == (1, 2, 4, 7, 15, 30)

    now = datetime(2026, 1, 1, 9, 0, 0)
    item = ReviewItem(id="t1", task_id="task", title="物理层", stage=0, status="active", next_review_at=now)
    assert next_interval_days(item) == 1

    expect_days = [2, 4, 7, 15, 30]
    for i, days in enumerate(expect_days, start=1):
        advance(item, now)
        assert item.stage == i, f"stage should be {i}"
        assert item.next_review_at == now + timedelta(days=days), f"next should be +{days}d"
        assert item.last_reviewed_at == now

    assert item.stage + 1 == MAX_STAGE
    advance(item, now)  # 第 6 次完成 → 已掌握
    assert item.stage == MAX_STAGE
    assert item.status == "mastered"


def test_review_retention_decays():
    """记忆保持率随间隔衰减：刚复习后接近 100%，超过间隔后显著下降。"""
    from datetime import datetime, timedelta

    from app.models import ReviewItem
    from app.services.review import retention_percent

    base = datetime(2026, 1, 1, 9, 0, 0)
    item = ReviewItem(
        id="t2", task_id="task", title="章节", stage=0, status="active",
        created_at=base, last_reviewed_at=base, next_review_at=base + timedelta(days=1),
    )
    fresh = retention_percent(item, base)
    after_1d = retention_percent(item, base + timedelta(days=1))
    after_3d = retention_percent(item, base + timedelta(days=3))
    assert fresh == 100
    assert after_1d < fresh
    assert after_3d < after_1d


def test_review_api_flow():
    """接口链路：上传生成复习项 → 概览 → 重置后成为待复习 → 完成一次推进 → 删除任务联动清理。"""
    from app.services.review import MAX_STAGE

    resp = client.post(
        "/api/upload",
        files={"file": ("review_demo.txt", SAMPLE_TEXT.encode("utf-8"), "text/plain")},
    )
    assert resp.status_code == 200, resp.text
    task = resp.json()
    task_id = task["id"]
    assert task["review_items"] > 0, "上传后应生成复习项"

    # 该任务的复习项
    items = client.get(f"/api/reviews?task_id={task_id}").json()
    assert len(items) == task["review_items"]
    assert all(it["stage"] == 0 and it["status"] == "active" for it in items)
    # 首次复习安排在 1 天后 → 此刻不应出现在待复习列表
    assert all(not it["is_due"] for it in items)

    # 概览里能看到这些条目
    summary = client.get("/api/reviews/summary").json()
    assert summary["total"] >= len(items)
    assert summary["intervals"] == [1, 2, 4, 7, 15, 30]

    # 重置其中一条 → 立即变为待复习
    first = items[0]["id"]
    client.post(f"/api/reviews/{first}/reset")
    due = client.get("/api/reviews/due").json()
    assert any(d["id"] == first for d in due)
    assert all(d["is_due"] for d in due)

    # 完成一次复习：stage 0 → 1，下次 +2 天
    done = client.post(f"/api/reviews/{first}/complete").json()
    assert done["stage"] == 1
    assert done["is_due"] is False
    assert "下次复习" in done["message"]
    assert done["max_stage"] == MAX_STAGE

    # 未来安排的日期分组可查询
    upcoming = client.get("/api/reviews/upcoming?days=7").json()
    assert isinstance(upcoming, list)

    # 删除任务应联动清理复习项
    client.delete(f"/api/tasks/{task_id}")
    assert client.get(f"/api/reviews?task_id={task_id}").json() == []


def test_review_progress_and_completion():
    """完成进度与一键完成：单条 master、按课件 complete-all、progress 统计。"""
    from app.services.review import MAX_STAGE

    resp = client.post(
        "/api/upload",
        files={"file": ("progress_demo.txt", SAMPLE_TEXT.encode("utf-8"), "text/plain")},
    )
    assert resp.status_code == 200, resp.text
    task_id = resp.json()["id"]
    items = client.get(f"/api/reviews?task_id={task_id}").json()
    assert len(items) >= 2

    # 该课件的进度行：初始 0%
    rows = client.get("/api/reviews/progress").json()
    row = next(r for r in rows if r["task_id"] == task_id)
    assert row["total"] == len(items)
    assert row["mastered"] == 0
    assert row["percent"] == 0
    assert row["stage_percent"] == 0
    assert row["stages_total"] == len(items) * MAX_STAGE

    # 单条「完成」→ 直接标记已掌握，stage 到顶
    first = items[0]["id"]
    done = client.post(f"/api/reviews/{first}/master").json()
    assert done["stage"] == MAX_STAGE
    assert done["status"] == "mastered"
    assert done["is_due"] is False
    assert "复习计划已完成" in done["message"]

    # 进度随之更新
    row = next(r for r in client.get("/api/reviews/progress").json() if r["task_id"] == task_id)
    assert row["mastered"] == 1
    assert row["stage_percent"] > 0
    assert row["percent"] == round(1 / len(items) * 100)

    # 一键完成该课件剩余计划
    result = client.post(f"/api/reviews/complete-all?task_id={task_id}").json()
    assert result["ok"] is True
    assert result["completed"] == len(items) - 1
    assert result["summary"]["total"] >= len(items)

    row = next(r for r in client.get("/api/reviews/progress").json() if r["task_id"] == task_id)
    assert row["mastered"] == row["total"]
    assert row["active"] == 0
    assert row["percent"] == 100
    assert row["due"] == 0

    # 已完成的项不再出现在待复习列表
    due_ids = {d["id"] for d in client.get("/api/reviews/due").json()}
    assert all(it["id"] not in due_ids for it in items)

    # 清理
    client.delete(f"/api/tasks/{task_id}")
    assert client.get(f"/api/reviews?task_id={task_id}").json() == []


def test_review_summary_progress_fields():
    """总览包含完成进度字段，且数值自洽。"""
    s = client.get("/api/reviews/summary").json()
    for key in ("total", "active", "mastered", "due", "stages_done", "stages_total", "progress_percent", "stage_percent"):
        assert key in s, f"缺少字段 {key}"
    assert s["total"] == s["active"] + s["mastered"]
    assert 0 <= s["progress_percent"] <= 100
    assert 0 <= s["stage_percent"] <= 100


def test_auth_requires_login():
    """未登录时：健康检查可用，业务接口一律 401。"""
    anon = TestClient(app)
    assert anon.get("/api/health").status_code == 200
    assert anon.get("/api/tasks").status_code == 401
    assert anon.get("/api/reviews/summary").status_code == 401
    assert anon.get("/api/settings").status_code == 401
    assert anon.post("/api/auth/logout").status_code == 200  # 幂等，未登录也能登出


def test_auth_register_login_logout():
    """注册 → 当前用户 → 登出 → 需要重新登录 → 再登录成功。"""
    c = TestClient(app)
    resp = c.post("/api/auth/register", json={"username": "ringo", "password": "secret123", "display_name": "Ringo"})
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["username"] == "ringo"
    assert data["display_name"] == "Ringo"
    assert "created_at" in data

    me = c.get("/api/auth/me")
    assert me.status_code == 200 and me.json()["username"] == "ringo"

    assert c.post("/api/auth/logout").status_code == 200
    assert c.get("/api/auth/me").status_code == 401
    assert c.get("/api/tasks").status_code == 401

    again = c.post("/api/auth/login", json={"username": "ringo", "password": "secret123"})
    assert again.status_code == 200, again.text
    assert c.get("/api/auth/me").status_code == 200
    # 密码错误
    assert c.post("/api/auth/login", json={"username": "ringo", "password": "wrong"}).status_code == 401


def test_auth_validation_rules():
    """用户名重复、密码过短、用户名非法字符均返回 400。"""
    c = TestClient(app)
    assert c.post("/api/auth/register", json={"username": "tester", "password": "whatever123"}).status_code == 400
    assert c.post("/api/auth/register", json={"username": "shortpw", "password": "123"}).status_code == 400
    assert c.post("/api/auth/register", json={"username": "bad name!", "password": "goodpass123"}).status_code == 400


def test_password_not_stored_in_plaintext():
    """口令必须以 PBKDF2 哈希存储，且同一口令两次哈希不同（随机盐）。"""
    from app.services.auth import hash_password, verify_password

    h1 = hash_password("hello12345")
    h2 = hash_password("hello12345")
    assert h1 != h2                      # 随机盐 → 哈希不同
    assert h1.startswith("pbkdf2_sha256$")
    assert "hello12345" not in h1        # 不含明文
    assert verify_password("hello12345", h1)
    assert not verify_password("hello12346", h1)
    assert not verify_password("hello12345", "garbage")


def test_user_data_isolation():
    """数据按用户隔离：A 上传的内容 B 看不到、也不能操作。"""
    a = make_client("iso_a", "passwordA1")
    b = make_client("iso_b", "passwordB1")

    resp = a.post("/api/upload", files={"file": ("a_only.txt", SAMPLE_TEXT.encode("utf-8"), "text/plain")})
    assert resp.status_code == 200, resp.text
    task = resp.json()
    task_id = task["id"]

    # A 能看到自己的任务
    a_tasks = a.get("/api/tasks").json()
    assert any(t["id"] == task_id for t in a_tasks)

    # B 看不到，也读不到 / 删不掉
    b_tasks = b.get("/api/tasks").json()
    assert all(t["id"] != task_id for t in b_tasks)
    assert b.get(f"/api/tasks/{task_id}").status_code == 404
    assert b.delete(f"/api/tasks/{task_id}").status_code == 404

    # B 的复习项里没有 A 的内容
    b_reviews = b.get("/api/reviews").json()
    assert all(r["task_id"] != task_id for r in b_reviews)
    assert all(r["filename"] != "a_only.txt" for r in b_reviews)

    # B 不能操作 A 的复习项
    a_items = a.get(f"/api/reviews?task_id={task_id}").json()
    assert a_items, "A 应该有自己的复习项"
    victim = a_items[0]["id"]
    assert b.post(f"/api/reviews/{victim}/complete").status_code == 404
    assert b.post(f"/api/reviews/{victim}/master").status_code == 404
    assert b.delete(f"/api/reviews/{victim}").status_code == 404

    # B 一键完成也只能作用于自己的数据
    b.post("/api/reviews/complete-all").json()
    assert a.get(f"/api/reviews?task_id={task_id}").json()[0]["status"] == "active"

    # 清理
    assert a.delete(f"/api/tasks/{task_id}").status_code == 200


def test_claim_orphan_data_on_fresh_db():
    """历史遗留数据（无归属）只被系统里第一个账号认领一次。"""
    import uuid

    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    from app.models import Base, ReviewItem, Task
    from app.services import auth

    db_file = os.path.join(tempfile.gettempdir(), "ai_notes_claim_test.db")
    if os.path.exists(db_file):
        os.remove(db_file)
    engine = create_engine(f"sqlite:///{db_file}")
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()

    # 旧版本产生的、没有归属用户的数据
    db.add(Task(id=uuid.uuid4().hex, user_id=None, filename="legacy.txt", file_type="txt"))
    db.add(ReviewItem(id=uuid.uuid4().hex, user_id=None, task_id="legacy-task", title="旧知识点"))
    db.commit()

    # 第一个账号注册 → 认领
    first = auth.create_user(db, "firstuser", "password1")
    db.commit()
    assert auth.claim_orphan_data(db, first.id) == 2
    db.commit()
    assert db.query(Task).filter(Task.user_id == first.id).count() == 1
    assert db.query(ReviewItem).filter(ReviewItem.user_id == first.id).count() == 1

    # 第二个账号注册 → 不再认领（避免多人环境抢占他人数据）
    second = auth.create_user(db, "seconduser", "password2")
    db.commit()
    db.add(Task(id=uuid.uuid4().hex, user_id=None, filename="orphan2.txt", file_type="txt"))
    db.commit()
    assert auth.claim_orphan_data(db, second.id) == 0
    assert db.query(Task).filter(Task.user_id == second.id).count() == 0

    db.close()
    engine.dispose()


def test_bearer_token_auth():
    """手机 App 场景：不依赖 Cookie，改用 Authorization: Bearer 鉴权。"""
    c = TestClient(app)
    resp = c.post("/api/auth/register", json={"username": "mobile_user", "password": "mobile123"})
    assert resp.status_code == 200, resp.text
    token = resp.json().get("token")
    assert token, "登录/注册响应必须返回会话 token（供 App 保存）"

    # 全新客户端（无 Cookie），只带 Authorization 头
    anon = TestClient(app)
    assert anon.get("/api/tasks").status_code == 401
    headers = {"Authorization": f"Bearer {token}"}
    assert anon.get("/api/tasks", headers=headers).status_code == 200
    assert anon.get("/api/auth/me", headers=headers).json()["username"] == "mobile_user"

    # 上传与复习接口同样可用
    up = anon.post(
        "/api/upload",
        files={"file": ("mobile.txt", SAMPLE_TEXT.encode("utf-8"), "text/plain")},
        headers=headers,
    )
    assert up.status_code == 200, up.text
    task_id = up.json()["id"]
    assert anon.get(f"/api/reviews?task_id={task_id}", headers=headers).json()

    # 登出后 token 立即失效
    assert anon.post("/api/auth/logout", headers=headers).status_code == 200
    assert anon.get("/api/tasks", headers=headers).status_code == 401

    # 非法 token
    assert anon.get("/api/tasks", headers={"Authorization": "Bearer not-a-real-token"}).status_code == 401

    c.delete(f"/api/tasks/{task_id}")   # 用 Cookie 客户端清理


def _run_all():
    tests = [(n, fn) for n, fn in sorted(globals().items()) if n.startswith("test_") and callable(fn)]
    failed = 0
    for name, fn in tests:
        try:
            fn()
            print(f"  PASS  {name}")
        except Exception as exc:
            failed += 1
            print(f"  FAIL  {name}: {exc}")
    print(f"\n{len(tests) - failed}/{len(tests)} 通过")
    return failed


if __name__ == "__main__":
    sys.exit(1 if _run_all() else 0)
