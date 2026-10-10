"""解析服务：PDF / PPTX / 图片(OCR) / 文本 → 纯文本。

输出约定：
- PDF：逐页文本拼接；若文本层不可用（扫描件、或字体缺少 Unicode 映射导致提取出乱码），
  自动渲染页面并走 OCR 兜底
- PPTX：每页输出 "# 标题" + 正文行（便于兜底提取建树）
- 图片：OCR 文本（依赖 rapidocr，未安装时抛 OCR 不可用异常）
"""
from __future__ import annotations

import os
import tempfile

SUPPORTED_EXTS = {"pdf", "pptx", "txt", "md", "png", "jpg", "jpeg"}

#: 文本层可用性判定：有效字符占比下限与最短长度
MIN_MEANINGFUL_RATIO = 0.5
MIN_TEXT_CHARS = 30
#: PDF OCR 兜底渲染参数
OCR_PDF_DPI = 200
OCR_PDF_MAX_PAGES = 20


def parse_file(path: str, filename: str) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf":
        return parse_pdf(path)
    if ext == "pptx":
        return parse_pptx(path)
    if ext in ("txt", "md"):
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    if ext in ("png", "jpg", "jpeg"):
        return parse_image(path)
    raise ValueError(f"不支持的文件类型: .{ext}（支持 {sorted(SUPPORTED_EXTS)}）")


def _meaningful_ratio(text: str) -> float:
    """有效字符（字母/数字/汉字等 alnum）在非空白字符中的占比。"""
    chars = [c for c in text if not c.isspace()]
    if not chars:
        return 0.0
    good = sum(1 for c in chars if c.isalnum())
    return good / len(chars)


def _text_is_usable(text: str) -> bool:
    """判断提取出的文本是否可用于后续提取（过滤空文本与乱码）。"""
    stripped = text.strip()
    if len(stripped) < MIN_TEXT_CHARS:
        return False
    return _meaningful_ratio(stripped) >= MIN_MEANINGFUL_RATIO


def parse_pdf(path: str) -> str:
    try:
        import pymupdf as fitz
    except ImportError:  # 旧版本 PyMuPDF 兼容
        import fitz

    doc = fitz.open(path)
    try:
        parts = []
        for page in doc:
            parts.append(page.get_text())
        text = "\n".join(parts)
        if _text_is_usable(text):
            return text

        # 文本层不可用 → 渲染页面 OCR 兜底（扫描件 / 字形无 Unicode 映射）
        ocr_text = _ocr_pdf(doc)
        if _text_is_usable(ocr_text):
            return ocr_text

        if text.strip():
            raise ValueError(
                "PDF 文本层无法识别（字体缺少 Unicode 映射，OCR 兜底也未识别出内容）。"
                "请改用清晰截图上传，或用其他工具重新导出该 PDF。"
            )
        raise ValueError("PDF 未提取到文本（可能是扫描件/图片型 PDF，请转为图片或使用带文字的 PDF）")
    finally:
        doc.close()


def parse_pptx(path: str) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    lines = []
    for slide in prs.slides:
        title = None
        bullets = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            if shape == slide.shapes.title and title is None:
                title = text
            elif title is None and shape.has_text_frame and shape.text_frame.text.strip():
                # 无标题版式：第一段文字当标题
                title = text
            else:
                bullets.append(text)
        if title:
            lines.append(f"# {title}")  # 一级标题，兜底提取时作为根节点下的章节
        lines.extend(bullets)
    return "\n".join(lines)


_ocr_engine = None


def _get_ocr_engine():
    """惰性单例 OCR 引擎（避免每次请求重新加载模型）。"""
    global _ocr_engine
    if _ocr_engine is None:
        try:
            from rapidocr import RapidOCR  # 新版（3.x，onnxruntime 引擎）

            _ocr_engine = RapidOCR()
        except ImportError:
            from rapidocr_onnxruntime import RapidOCR  # 旧版包兼容

            _ocr_engine = RapidOCR()
    return _ocr_engine


def _run_ocr(engine, image) -> list[str]:
    """调用 OCR 引擎，兼容 rapidocr 3.x 返回对象与旧版 (result, elapse) 元组。"""
    output = engine(str(image))
    if output is None:
        return []
    txts = getattr(output, "txts", None)  # rapidocr 3.x 返回 RapidOCROutput 对象
    if txts:
        return [str(t) for t in txts if str(t).strip()]
    try:
        result, _ = output
    except (TypeError, ValueError):
        return []
    return [line[1] for line in result if len(line) > 1 and line[1]]


def _ocr_pdf(doc) -> str:
    """把 PDF 页面渲染为图片后 OCR；OCR 依赖缺失或无可用页面时返回空串。"""
    try:
        engine = _get_ocr_engine()
    except ImportError:
        return ""

    texts: list[str] = []
    with tempfile.TemporaryDirectory() as tmp:
        for index, page in enumerate(doc):
            if index >= OCR_PDF_MAX_PAGES:
                break
            pix = page.get_pixmap(dpi=OCR_PDF_DPI)
            img_path = os.path.join(tmp, f"page_{index}.png")
            pix.save(img_path)
            texts.extend(_run_ocr(engine, img_path))
    return "\n".join(texts)


def parse_image(path: str) -> str:
    """图片 OCR（中文）。依赖 rapidocr（onnxruntime），未安装时给出明确提示。"""
    try:
        engine = _get_ocr_engine()
    except ImportError:
        raise ValueError(
            "OCR 依赖未安装：pip install -r requirements-ocr.txt "
            "（rapidocr 3.x + onnxruntime，Python 3.13/3.14 可安装）"
        )
    return "\n".join(_run_ocr(engine, path))
